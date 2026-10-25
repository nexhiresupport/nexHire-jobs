"""End-to-end smoke test: build a small Excel file and verify it."""
from datetime import datetime, timezone, timedelta
from pathlib import Path
import tempfile

from openpyxl import load_workbook

from jobfeed.models import Job
from jobfeed.output import write_excel

IST = timezone(timedelta(hours=5, minutes=30))


def sample_job(n=1) -> Job:
    j = Job(source="greenhouse", source_id=str(n), title="Data Engineer",
            company=f"Company {n}", location="New York, NY",
            url=f"https://company{n}.greenhouse.io/jobs/{n}",
            description="Build data pipelines with Spark and Snowflake. 3+ years experience.",
            posted_at=datetime(2025, 1, 15, 10, 0, tzinfo=timezone.utc),
            employment_hint="fulltime", remote_hint="Remote", via_job_board=False,
            publisher="", description_complete=True)
    j.work_type = "remote"
    j.employment_type = "Full-time"
    j.experience_years = 3
    j.recruiter_email = ""
    j.sponsorship_note = ""
    j.flags = []
    j.link_status = "open (live company feed)"
    j.exclude_reason = ""
    return j


def sample_excluded() -> Job:
    j = sample_job(99)
    j.company = "Federal Corp"
    j.exclude_reason = "Security clearance: …active Top Secret…"
    return j


def test_excel_created():
    jobs = [sample_job(i) for i in range(1, 6)]
    excluded = [sample_excluded()]
    log_rows = [("Run time (UTC)", "2025-01-15 10:00:00"), ("Lookback hours", "25")]

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "jobs.xlsx"
        result = write_excel(path, jobs, excluded, log_rows, max_desc=5000)
        assert result.exists()
        wb = load_workbook(result)
        assert "Jobs" in wb.sheetnames
        assert "Excluded" in wb.sheetnames
        assert "Run Log" in wb.sheetnames

        ws = wb["Jobs"]
        # Header row + 5 job rows = 6 rows
        assert ws.max_row == 6
        # Check a few cell values
        assert ws["A1"].value == "Company"
        assert ws["A2"].value == "Company 1"
        assert ws["D2"].value == "remote"    # work type

        ex = wb["Excluded"]
        assert ex.max_row == 2
        assert ex["A2"].value == "Federal Corp"
        assert "clearance" in (ex["E2"].value or "").lower()


def test_excel_empty_jobs():
    """Should not crash with no jobs."""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "empty.xlsx"
        result = write_excel(path, [], [], [("Test", "ok")], max_desc=5000)
        assert result.exists()
        wb = load_workbook(result)
        assert wb["Jobs"].max_row == 1   # header only
