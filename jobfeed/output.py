"""Excel export (+ optional email)."""
from __future__ import annotations

import os
import re
import smtplib
from datetime import timezone, timedelta
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .filters import suggested_profile
from .models import Job

IST = timezone(timedelta(hours=5, minutes=30))
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")

JOB_COLUMNS = [
    ("Company", 24), ("Role", 38), ("Location", 26), ("Work Type", 12), ("Employment Type", 17),
    ("Careers Link", 45), ("Recruiter Email", 28), ("careersLink for resume prompt", 40),
    ("Experience Req (yrs)", 12), ("Profile (3yr/5yr)", 10), ("Posted (IST)", 17), ("Source", 14),
    ("Sponsorship", 22), ("Flags", 34), ("Link Status", 26), ("Job Description", 90),
]


def _clean(v, limit=None):
    if v is None:
        return ""
    s = ILLEGAL.sub("", str(v))
    return s[:limit] if limit else s


def _resume_link(job: Job) -> str:
    """Your prompt: C2C -> careers/job URL; W2 -> recruiter email (fall back to URL if no email)."""
    if job.employment_type == "W2 Contract" and job.recruiter_email:
        return job.recruiter_email
    return job.url


def _header(ws, cols):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 32


def write_excel(path: Path, jobs: list[Job], excluded: list[Job], log_rows: list[tuple], max_desc: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    _header(ws, JOB_COLUMNS)
    jobs = sorted(jobs, key=lambda j: (j.posted_at is None, -(j.posted_at.timestamp() if j.posted_at else 0)))
    for r, j in enumerate(jobs, 2):
        posted = j.posted_at.astimezone(IST).strftime("%Y-%m-%d %H:%M") if j.posted_at else ""
        row = [j.company, j.title, j.location, j.work_type, j.employment_type, j.url, j.recruiter_email,
               _resume_link(j), j.experience_years if j.experience_years is not None else "",
               suggested_profile(j.experience_years), posted, j.source + (f" / {j.publisher}" if j.publisher else ""),
               j.sponsorship_note, "; ".join(j.flags), j.link_status, _clean(j.description, max_desc)]
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=_clean(v) if not isinstance(v, int) else v)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
        if j.url:
            ws.cell(row=r, column=6).hyperlink = j.url
            ws.cell(row=r, column=6).font = Font(color="1155CC", underline="single")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(JOB_COLUMNS))}{max(1, len(jobs) + 1)}"

    ex = wb.create_sheet("Excluded")
    ex_cols = [("Company", 24), ("Role", 38), ("Location", 26), ("Source", 14), ("Reason", 80), ("Link", 50)]
    _header(ex, ex_cols)
    for r, j in enumerate(excluded, 2):
        for c, v in enumerate([j.company, j.title, j.location, j.source, j.exclude_reason, j.url], 1):
            ex.cell(row=r, column=c, value=_clean(v))
    ex.auto_filter.ref = f"A1:F{max(1, len(excluded) + 1)}"

    lg = wb.create_sheet("Run Log")
    _header(lg, [("Step", 34), ("Detail", 100)])
    for r, (k, v) in enumerate(log_rows, 2):
        lg.cell(row=r, column=1, value=_clean(k))
        lg.cell(row=r, column=2, value=_clean(v))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def send_email(path: Path, n_jobs: int, run_date: str) -> str:
    user, pw, to = os.getenv("SMTP_USER"), os.getenv("SMTP_PASS"), os.getenv("EMAIL_TO")
    if not (user and pw and to):
        return "skipped (SMTP_USER / SMTP_PASS / EMAIL_TO not set)"
    msg = EmailMessage()
    msg["Subject"] = f"Data Engineer jobs - {run_date} - {n_jobs} new"
    msg["From"], msg["To"] = user, to
    msg.set_content(f"{n_jobs} new matching jobs. Excel attached.\n")
    msg.add_attachment(path.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=path.name)
    host = os.getenv("SMTP_HOST", "smtp.gmail.com")
    with smtplib.SMTP_SSL(host, int(os.getenv("SMTP_PORT", "465")), timeout=60) as s:
        s.login(user, pw)
        s.send_message(msg)
    return f"sent to {to}"
